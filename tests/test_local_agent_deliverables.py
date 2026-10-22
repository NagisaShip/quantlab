"""本地 Agent 交付工具：授权目录写/读、目录创建、Markdown→docx、汇报 PPT。

运行环境没有 python-docx / python-pptx 时，导出用例 pytest.skip（不安装依赖）；
其余用例只依赖标准库与后端既有依赖。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from conftest import login
from fastapi.testclient import TestClient

from quant_agent.api.app import create_app
from quant_agent.api.deps import AppServices, reset_services, set_services
from quant_agent.platform import context as platform_context
from quant_agent.workflows import agent as workflow_agent
from quant_agent.workflows.contracts import WorkflowError

TOOL_SIGNATURES = (
    "local_write_text",
    "local_make_dir",
    "local_read_text",
    "local_export_docx",
    "local_export_pptx",
)


def test_document_preview_pdf_pagination_and_image_context(local_project):
    fitz = pytest.importorskip("fitz")
    from quant_agent.agents import platform_tools as PT
    from quant_agent.local.documents import preview_images

    assert PT.registry()["local_preview_document"].approval == ""

    c, _services, project, tmp, _me = local_project
    target = project / "真实页面.pdf"
    with fitz.open() as document:
        for index in range(3):
            page = document.new_page()
            page.insert_text((40, 40), f"Page {index + 1}")
        document.save(target)
    result = workflow_agent.local_preview_document(str(target), limit=2)
    assert result["total_pages"] == 3 and result["next_page"] == 3
    assert result["visual_review"] == "pending"
    images = preview_images("local_preview_document", result)
    assert len(images) == 2 and images[0].startswith("data:image/png;base64,")
    assert preview_images("unrelated_tool", result) == []
    assert c.get(result["pages"][0]["url"]).headers["content-type"] == "image/png"
    assert c.get("/api/local/files/download", params={"path": str(target)}).status_code == 200
    tail = c.post("/api/local/files/document-preview?start=3", json={"path": str(target)})
    assert tail.status_code == 200 and tail.json()["next_page"] is None
    outside = tmp / "outside.pdf"
    outside.write_bytes(target.read_bytes())
    with pytest.raises(WorkflowError, match="授权"):
        workflow_agent.local_preview_document(str(outside))
    with pytest.raises(WorkflowError):
        workflow_agent.local_preview_document(str(target), limit=100)
    with fitz.open() as document:
        document.new_page()
        replacement = project / "changed.pdf"
        document.save(replacement)
    target.write_bytes(replacement.read_bytes())
    changed = workflow_agent.local_preview_document(str(target))
    assert changed["source_hash"] != result["source_hash"] and changed["total_pages"] == 1


def test_document_preview_reports_missing_office(local_project, monkeypatch):
    pytest.importorskip("fitz")
    pytest.importorskip("docx")
    from quant_agent.local import documents

    _c, _services, project, _tmp, _me = local_project
    exported = workflow_agent.local_export_docx("# 报告\n\n正文", "report.docx")
    monkeypatch.setattr(documents, "office_binary", lambda: None)
    with pytest.raises(WorkflowError, match="尚未检查排版"):
        workflow_agent.local_preview_document(exported["path"])
    assert (project / "report.docx").is_file()


def test_ppt_long_content_is_paginated_and_word_styles(local_project):
    pytest.importorskip("docx")
    _c, _services, _project, _tmp, _me = local_project
    from docx import Document
    from docx.oxml.ns import qn

    normalized = workflow_agent._normalize_slides([{"title": "长内容", "bullets": [f"要点{i}" * 60 for i in range(5)]}])
    assert len(normalized) > 1
    assert all(len(slide["bullets"]) <= 4 for slide in normalized)
    result = workflow_agent.local_export_docx("正文 **加粗** 和 *斜体*", "styled.docx")
    document = Document(result["path"])
    assert any(run.bold and run.text == "加粗" for run in document.paragraphs[0].runs)
    assert document.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "Microsoft YaHei"
    assert document.sections[0].page_height > document.sections[0].page_width


def test_real_office_pages_and_model_image_roundtrip(local_project):
    pytest.importorskip("fitz")
    pytest.importorskip("docx")
    pytest.importorskip("pptx")
    from types import SimpleNamespace

    from quant_agent.agents.base import Agent
    from quant_agent.llm.base import Message, Tool, ToolCall
    from quant_agent.local.documents import office_binary

    if not office_binary():
        pytest.skip("LibreOffice unavailable; exercised by CI documents job")
    _c, _services, _project, _tmp, _me = local_project
    word = workflow_agent.local_export_docx(
        "# Layout\n\n中文排版检查 **重点**\n\n| A | B |\n|---|---|\n| 1 | 2 |", "layout.docx"
    )
    deck = workflow_agent.local_export_pptx(
        [{"title": "Layout", "bullets": ["中文排版检查", "Evidence"]}], "layout.pptx", "Preview"
    )
    for exported in (word, deck):
        result = workflow_agent.local_preview_document(exported["path"])
        assert result["pages"] and result["rendered"]
        assert "Layout" in "".join(page["text"] for page in result["pages"])

    class Router:
        calls = 0

        def chat(self, messages, **kwargs):
            self.calls += 1
            if self.calls == 1:
                return SimpleNamespace(
                    content="",
                    provider="test",
                    model="vision",
                    usage={},
                    tool_calls=[
                        ToolCall(id="preview", name="local_preview_document", arguments={"path": word["path"]})
                    ],
                )
            assert messages[-1].role == "user" and messages[-1].images
            assert messages[-1].to_dict()["content"][1]["type"] == "image_url"
            assert messages[-2].role == "tool"
            return SimpleNamespace(content="Reviewed", provider="test", model="vision", usage={}, tool_calls=[])

    agent = Agent(
        name="reviewer",
        role="report",
        router=Router(),
        context=SimpleNamespace(
            build=lambda *_a, **_k: [Message(role="user", content="Review")], snapshot_hash=lambda: "test"
        ),
        tools=[
            Tool(
                name="local_preview_document",
                description="preview",
                parameters={},
                handler=workflow_agent.local_preview_document,
            )
        ],
    )
    assert agent.run("Review").final == "Reviewed"


@pytest.fixture
def client(tmp_path):
    reset_services()
    services = AppServices.build(store_root=str(tmp_path))
    set_services(services)
    with TestClient(create_app()) as c:
        me = login(c)
        yield c, services, me
    reset_services()


@pytest.fixture
def local_project(tmp_path, monkeypatch):
    project = tmp_path / "交付物项目"
    project.mkdir()
    monkeypatch.setenv("QUANTLAB_MODE", "local")
    monkeypatch.setenv("QUANTLAB_PROJECT", str(project))
    monkeypatch.setenv("QUANTLAB_LIBRARY", str(tmp_path / "library"))
    monkeypatch.setenv("QUANTLAB_CREDENTIALS", str(tmp_path / "keys"))
    reset_services()
    services = AppServices.build(store_root=str(project / ".quantlab"))
    set_services(services)
    with TestClient(create_app()) as c:
        c.headers["X-Requested-With"] = "quantlab-test"
        me = c.get("/api/auth/me").json()
        with platform_context.bind(me["workspace"]["id"], user_id=me["user"]["id"], role="admin"):
            yield c, services, project, tmp_path, me
    reset_services()


def test_local_write_read_roundtrip_and_overwrite_protection(local_project):
    _c, _s, project, _tmp, _me = local_project
    text = "# 研究结论\n\n中文内容\n第二行\n"
    written = workflow_agent.local_write_text("交付物/notes/报告.md", text)
    target = Path(written["path"])
    assert target.is_file() and target.is_relative_to(project)
    assert written["bytes"] == len(text.encode("utf-8")) and written["chars"] == len(text)
    assert written["overwritten"] is False and written["encoding"] == "utf-8"
    assert target.read_text(encoding="utf-8") == text

    with pytest.raises(WorkflowError, match="已存在"):
        workflow_agent.local_write_text("交付物/notes/报告.md", "另一份")
    assert target.read_text(encoding="utf-8") == text

    replaced = workflow_agent.local_write_text("交付物/notes/报告.md", "覆盖后的内容", overwrite=True)
    assert replaced["overwritten"] is True
    assert target.read_text(encoding="utf-8") == "覆盖后的内容"

    read = workflow_agent.local_read_text("交付物/notes/报告.md")
    assert read["text"] == "覆盖后的内容" and read["truncated"] is False and read["next_offset"] is None
    assert read["bytes"] == target.stat().st_size


def test_append_long_report_then_export_docx_and_pptx(local_project):
    _c, _s, project, _tmp, _me = local_project
    first = "# 实验报告\n\n" + "第一段。" * 40
    second = "\n\n## 结果\n\n" + "第二段。" * 40
    workflow_agent.local_write_text("交付物/实验报告.md", first, overwrite=True)
    appended = workflow_agent.local_write_text("交付物/实验报告.md", second, append=True)
    assert appended["appended"] is True
    md = (project / "交付物/实验报告.md").read_text(encoding="utf-8")
    assert md.startswith(first) and md.endswith(second)

    workflow_agent.local_write_text("交付物/models.py", "class Model:\n    pass\n")
    code_out = workflow_agent.local_append_code("交付物/实验报告.md", "交付物/models.py", "十七、核心程序代码")
    assert code_out["appended_chars"] > 0
    assert "```python" in (project / "交付物/实验报告.md").read_text(encoding="utf-8")

    pytest.importorskip("docx")
    pytest.importorskip("pptx")
    docx_out = workflow_agent.local_export_docx_file("交付物/实验报告.md", "交付物/实验报告.docx", "实验报告")
    assert Path(docx_out["path"]).is_file() and docx_out["bytes"] > 0
    from docx import Document

    document = Document(docx_out["path"])
    assert any("实验报告" in p.text for p in document.paragraphs)

    slides = {
        "slides": [
            {"title": "研究问题", "bullets": ["风险识别", "供应链网络"], "notes": "开场"},
            {"title": "结果", "bullets": ["AUC 0.9043"], "notes": "T-HGNNs"},
        ]
    }
    workflow_agent.local_write_text("交付物/slides.json", __import__("json").dumps(slides, ensure_ascii=False))
    pptx_out = workflow_agent.local_export_pptx_file("交付物/slides.json", "交付物/汇报.pptx", "汇报")
    assert Path(pptx_out["path"]).is_file() and pptx_out["slides"] == 2


def test_local_read_text_paginates_by_chars(local_project):
    _c, _s, _project, _tmp, _me = local_project
    workflow_agent.local_write_text("page.txt", "0123456789abcdefghij")
    first = workflow_agent.local_read_text("page.txt", offset=0, limit=10)
    assert first["text"] == "0123456789" and first["truncated"] is True and first["next_offset"] == 10
    second = workflow_agent.local_read_text("page.txt", offset=10, limit=10)
    assert second["text"] == "abcdefghij" and second["truncated"] is False and second["next_offset"] is None

    with pytest.raises(WorkflowError, match="offset"):
        workflow_agent.local_read_text("page.txt", offset=-1)
    with pytest.raises(WorkflowError, match="limit"):
        workflow_agent.local_read_text("page.txt", limit=0)
    with pytest.raises(WorkflowError, match="不存在"):
        workflow_agent.local_read_text("missing.txt")


def test_local_read_text_rejects_binary_and_office_files(local_project):
    _c, _s, project, _tmp, _me = local_project
    (project / "作业.docx").write_bytes(b"PK\x03\x04fake docx body")
    (project / "报告.pdf").write_bytes(b"%PDF-1.7 fake")
    (project / "伪文本.txt").write_bytes(b"abc\x00def")
    for name in ("作业.docx", "报告.pdf", "伪文本.txt"):
        with pytest.raises(WorkflowError, match=r"二进制|办公"):
            workflow_agent.local_read_text(name)
    # 目录不能用 read_text 读
    workflow_agent.local_make_dir("输出目录")
    with pytest.raises(WorkflowError, match="不是文件"):
        workflow_agent.local_read_text("输出目录")


def test_local_write_rejects_unauthorized_private_and_escaping_paths(local_project):
    _c, _s, project, tmp, _me = local_project
    outside = tmp / "未授权目录"
    outside.mkdir()
    with pytest.raises(WorkflowError, match="授权"):
        workflow_agent.local_write_text(str(outside / "x.txt"), "x")
    with pytest.raises(WorkflowError, match="授权"):
        workflow_agent.local_write_text("../escape.txt", "x")
    with pytest.raises(WorkflowError, match="平台数据目录"):
        workflow_agent.local_write_text(str(project / ".quantlab" / "hack.txt"), "x")
    with pytest.raises(WorkflowError, match="平台数据目录"):
        workflow_agent.local_write_text(str(project / ".quantlab" / "project.json"), "{}")
    with pytest.raises(WorkflowError, match="平台数据目录"):
        workflow_agent.local_read_text(str(project / ".quantlab" / "project.json"))

    # 用户显式授权的可写目录可以写；只读授权只能读不能写。
    profile = _s.local
    external = tmp / "外部输出"
    external.mkdir()
    profile.grant(str(external), "write")
    allowed = workflow_agent.local_write_text(str(external / "成果.txt"), "已授权")
    assert Path(allowed["path"]).is_file()

    readonly = tmp / "只读资料"
    readonly.mkdir()
    (readonly / "资料.txt").write_text("只读内容", encoding="utf-8")
    profile.grant(str(readonly), "read")
    assert workflow_agent.local_read_text(str(readonly / "资料.txt"))["text"] == "只读内容"
    with pytest.raises(WorkflowError, match="授权"):
        workflow_agent.local_write_text(str(readonly / "新文件.txt"), "不能写")


def test_local_write_enforces_two_megabyte_limit(local_project):
    _c, _s, _project, _tmp, _me = local_project
    with pytest.raises(WorkflowError, match="2MB"):
        workflow_agent.local_write_text("big.txt", "a" * (workflow_agent.MAX_TEXT_BYTES + 1))


def test_local_make_dir_creates_nested_directories(local_project):
    _c, _s, project, tmp, _me = local_project
    created = workflow_agent.local_make_dir("交付物/报告/图片")
    assert created["created"] is True
    assert Path(created["path"]).is_dir() and Path(created["path"]).is_relative_to(project)
    assert workflow_agent.local_make_dir("交付物/报告/图片")["created"] is False

    workflow_agent.local_write_text("占位.txt", "x")
    with pytest.raises(WorkflowError, match="不是目录"):
        workflow_agent.local_make_dir("占位.txt")
    with pytest.raises(WorkflowError, match="授权"):
        workflow_agent.local_make_dir(str(tmp / "未授权" / "目录"))


def test_local_tools_refuse_account_mode(client):
    _c, _services, me = client
    with platform_context.bind(me["workspace"]["id"], user_id=me["user"]["id"], role="admin"):
        with pytest.raises(WorkflowError, match="仅本地项目可用"):
            workflow_agent.local_write_text("x.txt", "x")
        with pytest.raises(WorkflowError, match="仅本地项目可用"):
            workflow_agent.local_make_dir("目录")
        with pytest.raises(WorkflowError, match="仅本地项目可用"):
            workflow_agent.local_read_text("x.txt")
        with pytest.raises(WorkflowError, match="仅本地项目可用"):
            workflow_agent.local_export_docx("# 标题", "x.docx")
        with pytest.raises(WorkflowError, match="仅本地项目可用"):
            workflow_agent.local_export_pptx([{"title": "页"}], "x.pptx")


def test_local_tools_are_confirm_save_in_lab_and_gated(local_project):
    _c, _s, _project, _tmp, _me = local_project
    from quant_agent.agents import platform_tools as PT
    from quant_agent.agents.lab_state import new_project
    from quant_agent.api.routes_lab import _lab

    registry = PT.registry()
    for name in TOOL_SIGNATURES:
        tool = registry[name]
        assert tool.risk == "confirm", name
        assert tool.approval == "save", name
        assert "授权" in tool.description and "密钥" in tool.description, name
        assert tool.scene == "workflow"

    lab = _lab()
    project = lab.store.save(new_project({"idea": "交付物审批"}))
    # manual / accept_edits 会先弹审批卡片；auto 模式按 Lab 既有权限规则直接执行。
    assert lab._approval_category(project, "manual", "local_write_text") == "save"
    assert lab._approval_category(project, "accept_edits", "local_export_docx") == "save"
    assert lab._approval_category(project, "auto", "local_write_text") == ""

    # 从平台工具桥接层实际调用一次，确认注册的 handler 可写文件。
    bridge = lab._platform_bridge(project["id"], set())
    for name in TOOL_SIGNATURES:
        assert name in bridge
    result = bridge["local_write_text"].handler(path="桥接验证.txt", text="通过平台工具写入")
    assert not result.get("error"), result
    assert (Path(result["path"])).read_text(encoding="utf-8") == "通过平台工具写入"


def test_export_without_library_gives_actionable_install_hint(local_project, monkeypatch):
    _c, _s, project, _tmp, _me = local_project
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(WorkflowError, match="pip install"):
        workflow_agent.local_export_docx("# 标题\n正文", str(project / "无库.docx"))
    monkeypatch.setitem(sys.modules, "pptx", None)
    with pytest.raises(WorkflowError, match="pip install"):
        workflow_agent.local_export_pptx([{"title": "页", "bullets": ["要点"]}], str(project / "无库.pptx"))


def test_export_rejects_wrong_extension_and_unathorized_path(local_project):
    _c, _s, _project, tmp, _me = local_project
    with pytest.raises(WorkflowError, match=r"\.docx"):
        workflow_agent.local_export_docx("# 标题", "交付物/报告.md")
    with pytest.raises(WorkflowError, match=r"\.pptx"):
        workflow_agent.local_export_pptx([{"title": "页"}], "交付物/汇报.doc")
    with pytest.raises(WorkflowError, match="授权"):
        workflow_agent.local_export_docx("# 标题", str(tmp / "未授权" / "报告.docx"))
    with pytest.raises(WorkflowError, match="1—80"):
        workflow_agent.local_export_pptx([], "交付物/空.pptx")


def test_export_docx_from_markdown(local_project):
    docx = pytest.importorskip("docx", reason="环境未安装 python-docx（可选依赖 .[deliverables]），跳过导出用例")
    _c, _s, project, _tmp, _me = local_project
    markdown = (
        "# 作业2 研究报告\n\n"
        "正文第一段，含 **加粗** 与 `行内代码`。\n\n"
        "## 数据与方法\n\n"
        "- 要点一\n"
        "- 要点二\n\n"
        "1. 第一步\n"
        "2. 第二步\n\n"
        "| 指标 | 数值 |\n"
        "| --- | --- |\n"
        "| 夏普 | 1.2 |\n"
        "| 回撤 | -8% |\n\n"
        "```python\nprint('hi')\n```\n\n"
        "> 结论引用\n"
    )
    result = workflow_agent.local_export_docx(markdown, "交付物/作业2.docx", title="作业2 交付")
    path = Path(result["path"])
    assert path.is_file() and path.is_relative_to(project) and path.suffix == ".docx"
    assert result["bytes"] > 1000 and result["format"] == "docx" and result["library"] == "python-docx"
    assert result["tables"] == 1 and result["code_blocks"] == 1 and result["headings"] >= 2

    document = docx.Document(str(path))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "作业2 交付" in texts
    assert any("作业2 研究报告" in text for text in texts)
    assert any("正文第一段" in text and "**" not in text for text in texts)
    assert any("要点一" in text for text in texts)
    table = document.tables[0]
    assert table.rows[0].cells[0].text == "指标" and table.rows[1].cells[1].text == "1.2"

    again = workflow_agent.local_export_docx("# 新标题\n", "交付物/作业2.docx")
    assert again["overwritten"] is True and Path(again["path"]).is_file()


def test_export_pptx_with_notes_and_16_9(local_project):
    pptx = pytest.importorskip("pptx", reason="环境未安装 python-pptx（可选依赖 .[deliverables]），跳过导出用例")
    _c, _s, project, _tmp, _me = local_project
    slides = [
        {"title": "研究背景", "bullets": ["选题：量化因子", "目标：稳健收益"], "notes": "开场先讲背景"},
        {"title": "结论", "bullets": ["夏普 1.2", "最大回撤 8%"]},
    ]
    result = workflow_agent.local_export_pptx(slides, "交付物/答辩.pptx", title="作业2 答辩")
    path = Path(result["path"])
    assert path.is_file() and path.is_relative_to(project) and path.suffix == ".pptx"
    assert result["slides"] == 2 and result["aspect"] == "16:9" and result["library"] == "python-pptx"

    presentation = pptx.Presentation(str(path))
    assert len(presentation.slides) == 3  # 封面 + 两页
    assert presentation.slide_width > presentation.slide_height

    def slide_text(slide) -> str:
        return "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))

    assert "作业2 答辩" in slide_text(presentation.slides[0])
    assert "研究背景" in slide_text(presentation.slides[1])
    assert "选题：量化因子" in slide_text(presentation.slides[1]) and "目标：稳健收益" in slide_text(
        presentation.slides[1]
    )
    assert presentation.slides[1].notes_slide.notes_text_frame.text == "开场先讲背景"
    assert "结论" in slide_text(presentation.slides[2])

    again = workflow_agent.local_export_pptx(slides, "交付物/答辩.pptx")
    assert again["overwritten"] is True


def test_relative_paths_resolve_from_local_project_root(local_project):
    """Agent 常传项目内相对路径；必须相对项目根解析，不能落到后端进程 cwd。"""
    _c, services, project, _tmp, _me = local_project
    inputs = project / "验收输入"
    inputs.mkdir()
    summary = inputs / "关键证据摘要.md"
    summary.write_text("# 验收输入\n\n- 关键证据一\n- 关键证据二\n", encoding="utf-8")

    profile = services.local
    assert profile.authorize("验收输入") == inputs.resolve()
    assert profile.authorize(str(inputs)) == inputs.resolve()
    browsed = profile.browse("验收输入")
    assert Path(browsed["path"]) == inputs.resolve()
    assert any(entry["name"] == "关键证据摘要.md" for entry in browsed["entries"])

    read = workflow_agent.local_read_text("验收输入/关键证据摘要.md")
    assert "关键证据一" in read["text"]
    assert read["has_more"] is False and read["next_offset"] is None
    assert "已读到文件结尾" in read["more_hint"]
    assert Path(read["path"]) == summary.resolve()

    # workflow_local_files preview 对 .md 走文本读取，而不是“只支持表格”的报错。
    preview = workflow_agent.local_files("验收输入/关键证据摘要.md", action="preview")
    assert preview["kind"] == "text" and "关键证据一" in preview["text"]

    # 相对路径写文件、相对路径引用导入同样按项目根解析。
    written = workflow_agent.local_write_text("验收输入/新文件.md", "# 新文件\n")
    assert Path(written["path"]) == (inputs / "新文件.md").resolve()
    imported = profile.import_file(workflow_agent.service(), "验收输入/关键证据摘要.md", "reference")
    assert Path(imported["source"]) == summary.resolve()


def test_local_text_preview_supports_md_txt_json_py_and_pages(local_project):
    _c, _s, project, _tmp, _me = local_project
    (project / "任务说明.md").write_text("# 任务\n\n写一份报告。\n", encoding="utf-8")
    (project / "notes.txt").write_text("纯文本说明", encoding="utf-8")
    (project / "config.json").write_text('{"answer": 42}', encoding="utf-8")
    (project / "calc.py").write_text("print('hello')\n", encoding="utf-8")
    for name, needle in (
        ("任务说明.md", "写一份报告"),
        ("notes.txt", "纯文本说明"),
        ("config.json", '"answer": 42'),
        ("calc.py", "print('hello')"),
    ):
        preview = workflow_agent.local_files(name, action="preview")
        assert preview["kind"] == "text" and needle in preview["text"], name

    # 默认一次返回 20k 字符；超过时明确“还有更多”并给出 offset 分段方式。
    (project / "长文档.md").write_text("甲" * 25000, encoding="utf-8")
    first = workflow_agent.local_files("长文档.md", action="preview")
    assert first["truncated"] is True and first["has_more"] is True
    assert first["next_offset"] == workflow_agent.TEXT_READ_LIMIT
    assert "还有更多" in first["more_hint"]
    tail = workflow_agent.local_files("长文档.md", action="preview", offset=first["next_offset"])
    assert len(tail["text"]) == 5000 and tail["truncated"] is False


def test_local_preview_binary_message_is_actionable(local_project):
    _c, _s, project, _tmp, _me = local_project
    (project / "作业模板.docx").write_bytes(b"PK\x03\x04fake")
    (project / "图表.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    for name in ("作业模板.docx", "图表.png"):
        with pytest.raises(WorkflowError, match=r"二进制|办公") as excinfo:
            workflow_agent.local_files(name, action="preview")
        message = str(excinfo.value)
        assert "local_read_text" in message and "系统应用" in message


def test_unauthorized_error_suggests_absolute_path_or_grant(local_project):
    _c, _s, _project, tmp, _me = local_project
    outside = tmp / "未授权资料"
    outside.mkdir()
    target = outside / "说明.md"
    target.write_text("未授权内容", encoding="utf-8")

    with pytest.raises(WorkflowError, match="绝对路径") as read_error:
        workflow_agent.local_read_text(str(target))
    assert "local_read_text" in str(read_error.value)

    with pytest.raises(WorkflowError, match="绝对路径") as write_error:
        workflow_agent.local_write_text(str(outside / "新文件.txt"), "x")
    assert "local_write_text" in str(write_error.value)
    assert "文件与目录权限" in str(write_error.value)


def test_local_preview_route_reads_markdown_and_relative_path(local_project):
    c, _s, project, _tmp, _me = local_project
    inputs = project / "验收输入"
    inputs.mkdir()
    (inputs / "任务说明.md").write_text("# 任务\n\n请产出报告。\n", encoding="utf-8")

    browsed = c.get("/api/local/files", params={"path": "验收输入"})
    assert browsed.status_code == 200, browsed.text
    assert Path(browsed.json()["path"]) == inputs.resolve()

    response = c.get("/api/local/files/preview", params={"path": "验收输入/任务说明.md"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["kind"] == "text" and "请产出报告" in body["text"]

    (project / "模板.docx").write_bytes(b"PK\x03\x04fake")
    binary = c.get("/api/local/files/preview", params={"path": "模板.docx"})
    assert binary.status_code == 400, binary.text
    assert "二进制" in binary.text and "系统应用" in binary.text


def test_local_chart_and_embeds_into_docx_pptx(local_project):
    _c, _s, _project, _tmp, _me = local_project
    chart = workflow_agent.local_chart("bar", "风险分布", ["高", "中", "低"], [12, 30, 58], "交付物/图表.png")
    assert chart["format"] == "png" and Path(chart["path"]).is_file()
    assert chart["url"].startswith("/api/local/files/image?path=")
    image = _c.get("/api/local/files/image", params={"path": "交付物/图表.png"})
    assert image.status_code == 200 and image.headers["content-type"] == "image/png"

    md = "# 报告\n\n下面是图表：\n\n![风险分布](交付物/图表.png)\n"
    docx_out = workflow_agent.local_export_docx(md, "交付物/带图报告.docx", "带图报告")
    assert docx_out["images"] == 1

    slides = [
        {"layout": "section", "title": "研究结论", "bullets": ["数据与口径", "结果与限制"]},
        {"layout": "image", "title": "风险分布", "bullets": ["高/中/低三档"], "image": "交付物/图表.png"},
    ]
    pptx_out = workflow_agent.local_export_pptx(slides, "交付物/带图汇报.pptx", "带图汇报")
    assert pptx_out["slides"] == 2 and Path(pptx_out["path"]).stat().st_size > 0
    from pptx import Presentation

    deck = Presentation(pptx_out["path"])
    assert len(deck.slides) == 3
    assert any(shape.shape_type == 13 for shape in deck.slides[2].shapes)  # 13 = PICTURE
