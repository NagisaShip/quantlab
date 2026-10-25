# ruff: noqa: F811
"""Actual embedded media/data and honest page counts, through public exporters."""

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from test_local_agent_deliverables import local_project  # noqa: F401

from quant_agent.local.document_charts import normalize_chart
from quant_agent.workflows import agent
from quant_agent.workflows.contracts import WorkflowError

CHART = {
    "kind": "bar",
    "title": "实际收益对比",
    "labels": ["策略甲", "策略乙"],
    "series": [{"name": "样本内", "values": [-2.5, 4.0]}, {"name": "样本外", "values": [-1.0, 2.0]}],
    "unit": "%",
    "source": "fixture-only",
}


def test_word_contains_chart_media_caption_and_not_json_code(local_project):
    from docx import Document

    _, _, _, _, _ = local_project
    result = agent.local_export_docx(
        "## 观测结果\n\n```chart\n" + json.dumps(CHART, ensure_ascii=False) + "\n```\n\n结论。", "含图.docx"
    )
    document = Document(result["path"])
    assert result["charts"] == result["images"] == len(document.inline_shapes) == 1
    text = "\n".join(p.text for p in document.paragraphs)
    assert "单位：%" in text and "来源：fixture-only" in text and '"series"' not in text
    with ZipFile(result["path"]) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_pptx_native_charts_keep_values_and_editable_workbook(local_project, kind):
    from pptx import Presentation

    chart = {**CHART, "kind": kind}
    if kind == "pie":
        chart["series"] = [{"name": "占比", "values": [25, 75]}]
    result = agent.local_export_pptx(
        [{"title": "证据", "chart": chart}], f"{kind}.pptx", include_cover=False, expected_slides=1
    )
    deck = Presentation(result["path"])
    native = next(shape.chart for shape in deck.slides[0].shapes if shape.has_chart)
    assert list(native.series[0].values) == chart["series"][0]["values"]
    if kind == "bar":
        assert all(not series.invert_if_negative for series in native.series)
    assert result["slides"] == len(deck.slides) == 1 and result["charts"] == 1
    with ZipFile(result["path"]) as archive:
        assert any(name.startswith("ppt/charts/chart") for name in archive.namelist())
        assert any(name.startswith("ppt/embeddings/") and name.endswith(".xlsx") for name in archive.namelist())


def test_page_budget_refuses_extra_cover_or_overflow_without_overwriting(local_project):
    _, _, project, _, _ = local_project
    result = agent.local_export_pptx([{"title": "原页"}], "count.pptx", include_cover=False, expected_slides=1)
    original = Path(result["path"]).read_bytes()
    with pytest.raises(WorkflowError, match="实际共2页"):
        agent.local_export_pptx([{"title": "结果"}], "count.pptx", title="封面", expected_slides=1)
    assert (project / "count.pptx").read_bytes() == original
    assert agent.local_export_pptx([{"title": "结果"}], "no-cover.pptx")["slides"] == 1


def test_json_file_keeps_chart_and_total_page_policy(local_project):
    plan = {
        "title": "测试",
        "include_cover": False,
        "expected_slides": 1,
        "slides": [{"title": "比较", "chart": CHART}],
    }
    agent.local_write_text("plan.json", json.dumps(plan, ensure_ascii=False))
    result = agent.local_export_pptx_file("plan.json", "plan.pptx")
    assert result["slides"] == result["charts"] == 1


def test_image_paths_with_spaces_embed_and_missing_assets_fail(local_project):
    from docx import Document

    image = agent.local_chart("line", "变化", ["甲", "乙"], [1, 2], "figures/chart space.png")
    out = agent.local_export_docx(f"![变化](<{image['path']}>)", "space.docx")
    assert len(Document(out["path"]).inline_shapes) == 1
    inline = agent.local_export_docx(f"图前说明 ![变化](<{image['path']}>) 图后结论", "inline.docx")
    parsed = Document(inline["path"])
    assert len(parsed.inline_shapes) == 1
    assert "图前说明" in "\n".join(p.text for p in parsed.paragraphs)
    assert "图后结论" in "\n".join(p.text for p in parsed.paragraphs)
    with pytest.raises(WorkflowError, match="图片未找到"):
        agent.local_export_docx("![缺失](not-here.png)", "missing.docx")
    with pytest.raises(WorkflowError, match="图片未找到"):
        agent.local_export_pptx([{"title": "缺失", "image": "not-here.png"}], "missing.pptx")


@pytest.mark.parametrize("values", [[True, 2], [float("nan"), 2], [float("inf"), 2], [None, 2], [1]])
def test_invalid_chart_values_are_not_silently_replaced(values):
    with pytest.raises(WorkflowError):
        normalize_chart({**CHART, "series": [{"name": "bad", "values": values}]})


def test_negative_pie_and_conflicting_assets_are_rejected():
    with pytest.raises(WorkflowError, match="非负"):
        normalize_chart({**CHART, "kind": "pie", "series": CHART["series"][:1]})
    with pytest.raises(WorkflowError, match="同时"):
        agent._normalize_slides([{"title": "冲突", "chart": CHART, "image": "chart.png"}])


def test_chart_block_and_slide_cannot_succeed_without_chart_data(local_project):
    with pytest.raises(WorkflowError, match="结束围栏"):
        agent.local_export_docx("```chart\n" + json.dumps(CHART), "broken.docx")
    with pytest.raises(WorkflowError, match="缺少 chart"):
        agent.local_export_pptx([{"title": "没有数据", "layout": "chart"}], "broken.pptx")


def test_office_preview_survives_invalid_temp_and_unicode_source(local_project, monkeypatch):
    from quant_agent.local.documents import office_binary

    if not office_binary():
        pytest.skip("LibreOffice unavailable")
    monkeypatch.setenv("TEMP", "relative/nonexistent")
    monkeypatch.setenv("TMP", "relative/nonexistent")
    out = agent.local_export_docx(
        "图表来自固定测试数据。\n\n```chart\n" + json.dumps(CHART) + "\n```", "中文 含图.docx"
    )
    preview = agent.local_preview_document(out["path"])
    assert preview["rendered"] and preview["total_pages"] == 1
    assert "fixture-only" in preview["pages"][0]["text"]
