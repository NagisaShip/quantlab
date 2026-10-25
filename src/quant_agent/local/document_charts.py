"""Bounded, shared chart data for Word figures and editable PowerPoint charts."""

from __future__ import annotations

import io
import math
from pathlib import Path

from ..workflows.contracts import WorkflowError


def normalize_chart(raw):
    if not isinstance(raw, dict):
        raise WorkflowError("chart 须为对象")
    kind = raw.get("kind", "bar")
    labels = raw.get("labels")
    series = raw.get("series")
    if series is None and "values" in raw:
        series = [{"name": raw.get("name", "数值"), "values": raw["values"]}]
    if kind not in {"bar", "line", "pie"}:
        raise WorkflowError("图表 kind 只支持 bar / line / pie")
    if not isinstance(labels, list) or not 2 <= len(labels) <= 40:
        raise WorkflowError("图表 labels 须为 2—40 个类别")
    if kind == "pie" and len(labels) > 12:
        raise WorkflowError("饼图最多支持12个类别，请合并已知类别或改用柱图")
    if any(not isinstance(label, str) or not label.strip() or len(label) > 40 for label in labels):
        raise WorkflowError("图表类别须为 1—40 字的非空字符串")
    if not isinstance(series, list) or not 1 <= len(series) <= 6 or (kind == "pie" and len(series) != 1):
        raise WorkflowError("图表须有 1—6 组 series，饼图只支持一组")
    clean = []
    for item in series:
        if not isinstance(item, dict):
            raise WorkflowError("series 每组须为 {name, values}")
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(labels):
            raise WorkflowError("图表 values 与 labels 须等长")
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
            raise WorkflowError("图表 values 须为有限数字，不能包含布尔值、NaN/Inf 或缺失值")
        if kind == "pie" and (min(values) < 0 or sum(values) <= 0):
            raise WorkflowError("饼图数值须非负且合计大于零")
        clean.append({"name": str(item.get("name", "数值"))[:60], "values": values})
    return {
        "kind": kind,
        "labels": labels,
        "series": clean,
        **{key: str(raw.get(key) or "")[:300] for key in ("title", "unit", "source")},
    }


def chart_caption(chart):
    return "；".join(
        value
        for value in [
            chart["title"],
            "单位：" + chart["unit"] if chart["unit"] else "",
            "来源：" + chart["source"] if chart["source"] else "",
        ]
        if value
    )


def chart_png(chart):
    """Render only supplied values; retain signs, labels, units and provenance."""
    from PIL import Image, ImageDraw, ImageFont

    def font(size):
        for name in [
            "msyh.ttc",
            "C:/Windows/Fonts/msyh.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
            "/System/Library/Fonts/PingFang.ttc",
            "DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        ]:
            try:
                return ImageFont.truetype(str(Path(name)), size)
            except OSError:
                continue
        return ImageFont.load_default()

    # Extra rows accommodate all labels rather than silently dropping categories.
    n = len(chart["labels"])
    w, h = max(1200, n * 145), 850
    image = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(image)
    small = font(20)
    colors = ["#15803D", "#0284C7", "#B45309", "#BE123C", "#7C3AED", "#0F766E"]
    draw.text((60, 20), chart["title"], font=font(34), fill="#0F172A")
    draw.text((60, 66), "单位：" + chart["unit"] if chart["unit"] else "", font=small, fill="#475569")
    for index, series in enumerate(chart["series"]):
        draw.text((60 + index * 190, 105), series["name"][:12], font=small, fill=colors[index])
    left, top, right, bottom = 120, 155, w - 55, 570
    if chart["kind"] == "pie":
        values = chart["series"][0]["values"]
        total, angle = sum(values), -90
        for index, (label, value) in enumerate(zip(chart["labels"], values, strict=True)):
            end = angle + value / total * 360
            draw.pieslice((left, top, left + 400, top + 400), angle, end, fill=colors[index % len(colors)])
            row, col = index % 12, index // 12
            draw.text(
                (570 + col * 380, 155 + row * 34), f"{label}: {value:g}", font=small, fill=colors[index % len(colors)]
            )
            angle = end
    else:
        values = [v for s in chart["series"] for v in s["values"]]
        low, high = min(0, min(values)), max(0, max(values))
        if chart["kind"] == "line":
            padding = (max(values) - min(values)) * 0.1 or max(abs(max(values)) * 0.05, 1)
            low, high = min(values) - padding, max(values) + padding
        span = high - low or 1

        def y(value):
            return bottom - (value - low) / span * (bottom - top)

        for tick in range(6):
            value = low + span * tick / 5
            py = y(value)
            draw.line((left, py, right, py), fill="#E2E8F0", width=2)
            draw.text((10, py - 12), f"{value:.4g}", font=small, fill="#475569")
        step = (right - left) / n
        for si, series in enumerate(chart["series"]):
            points = []
            for index, value in enumerate(series["values"]):
                x = left + step * (index + 0.5)
                py = y(value)
                if chart["kind"] == "bar":
                    bw = step * 0.75 / len(chart["series"])
                    bx = left + step * index + step * 0.125 + si * bw
                    draw.rectangle((bx, min(y(0), py), bx + bw - 2, max(y(0), py)), fill=colors[si])
                else:
                    points.append((x, py))
                    draw.ellipse((x - 5, py - 5, x + 5, py + 5), fill=colors[si])
            if len(points) > 1:
                draw.line(points, fill=colors[si], width=4)
        for index, label in enumerate(chart["labels"]):
            # Wrap full category labels instead of silently dropping them.
            tile = Image.new("RGB", (max(40, int(step) - 6), 180), "white")
            td = ImageDraw.Draw(tile)
            for row in range(0, min(len(label), 40), 6):
                td.text((0, row // 6 * 22), label[row : row + 6], font=small, fill="#334155")
            image.paste(tile, (int(left + step * index), bottom + 10))
    draw.text((60, h - 35), "来源：" + chart["source"] if chart["source"] else "", font=small, fill="#475569")
    if not chart["title"] and not chart["source"]:
        # Word supplies a real caption outside the bitmap. Do not reserve blank
        # title/source bands that force an otherwise fitting figure onto a new page.
        label_rows = max((len(label) + 5) // 6 for label in chart["labels"])
        image = image.crop((0, 60, w, min(h, bottom + max(44, label_rows * 22) + 30)))
    output = io.BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


def add_pptx_chart(slide, chart, left, top, width, height):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
    from pptx.util import Pt

    data = CategoryChartData()
    data.categories = chart["labels"]
    for series in chart["series"]:
        data.add_series(series["name"], series["values"])
    types = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}
    native = slide.shapes.add_chart(types[chart["kind"]], left, top, width, height, data).chart
    native.has_title = False
    native.font.name, native.font.size = "Microsoft YaHei", Pt(14)
    native.has_legend = len(chart["series"]) > 1 or chart["kind"] == "pie"
    if native.has_legend:
        native.legend.position = XL_LEGEND_POSITION.BOTTOM
        native.legend.include_in_layout = False
    if chart["kind"] != "pie":
        native.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        if chart["kind"] == "bar":
            for series in native.series:
                series.invert_if_negative = False
        native.value_axis.has_title = bool(chart["unit"])
        if chart["unit"]:
            native.value_axis.axis_title.text_frame.text = chart["unit"]
    else:
        native.plots[0].has_data_labels = True
        native.plots[0].data_labels.show_value = True
    return native
